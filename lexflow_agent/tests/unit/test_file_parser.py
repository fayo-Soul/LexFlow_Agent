"""文件解析工具单元测试

注意: MinerU 相关测试需要 conda 环境，CI 环境中跳过。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from lexflow_agent.engine.tools.file_parser import (
    SUPPORTED_EXTENSIONS,
    FileParser,
    FileParserError,
    _generate_document_id,
    _text_to_pages,
    _validate_file,
    _validate_file_signature,
)


# ── 文档 ID 生成 ──────────────────────────────────────────

def test_generate_document_id_stable():
    """相同内容和文件名生成相同 ID"""
    content = b"hello world"
    id1 = _generate_document_id("劳动合同.pdf", content)
    id2 = _generate_document_id("劳动合同.pdf", content)
    assert id1 == id2
    assert id1.startswith("doc_")


def test_generate_document_id_different_content():
    """不同内容生成不同 ID"""
    id1 = _generate_document_id("a.txt", b"content1")
    id2 = _generate_document_id("a.txt", b"content2")
    assert id1 != id2


# ── 文件校验 ──────────────────────────────────────────────

def test_validate_unsupported_extension():
    with pytest.raises(FileParserError, match="不支持的文件格式"):
        _validate_file(".exe", b"dummy")


def test_validate_file_too_large():
    with pytest.raises(FileParserError, match="文件过大"):
        _validate_file(".txt", b"x" * (51 * 1024 * 1024))


def test_validate_all_supported_extensions():
    for ext in SUPPORTED_EXTENSIONS:
        _validate_file(ext, b"x" * 100)  # 不抛异常


def test_validate_file_signature_pdf():
    _validate_file_signature(".pdf", b"%PDF-1.4 ...")
    with pytest.raises(FileParserError, match="文件头与扩展名不匹配"):
        _validate_file_signature(".pdf", b"Not a PDF")


def test_validate_file_signature_docx():
    _validate_file_signature(".docx", b"PK\x03\x04...")
    with pytest.raises(FileParserError, match="文件头与扩展名不匹配"):
        _validate_file_signature(".docx", b"Not a DOCX")


def test_validate_file_signature_skip_unsupported():
    """TXT/MD 无头校验，不应抛异常"""
    _validate_file_signature(".txt", b"anything")
    _validate_file_signature(".md", b"anything")


# ── 文本分页 ──────────────────────────────────────────────

def test_text_to_pages_empty():
    pages = _text_to_pages("")
    assert len(pages) == 1
    assert "无文本内容" in pages[0].text


def test_text_to_pages_single_page():
    text = "第一段\n\n第二段\n\n第三段"
    pages = _text_to_pages(text, max_chars=1000)
    assert len(pages) == 1


def test_text_to_pages_multi_page():
    """短 max_chars 应触发多页"""
    para = "A" * 100
    text = "\n\n".join([para] * 20)
    pages = _text_to_pages(text, max_chars=150)
    assert len(pages) > 1
    # 每页不应该截断段落
    for page in pages:
        assert page.page >= 1
        assert page.text
    # 总字符数不应丢失太多（允许少量截断）
    total = sum(len(p.text) for p in pages)
    original = len(text)
    assert total >= original * 0.8  # 不丢超过20%


def test_text_to_pages_page_numbering():
    text = "\n\n".join([f"段落{i}" for i in range(30)])
    pages = _text_to_pages(text, max_chars=50)
    for i, page in enumerate(pages, start=1):
        assert page.page == i


# ── FileParser ────────────────────────────────────────────

def test_parse_txt():
    parser = FileParser()
    content = "这是第一段文字。\n\n这是第二段文字。"
    doc = parser.parse_from_bytes("test.txt", content.encode("utf-8"))
    assert doc.document_id.startswith("doc_")
    assert doc.file_name == "test.txt"
    assert len(doc.pages) >= 1
    assert "第一段" in doc.pages[0].text


def test_parse_md():
    parser = FileParser()
    content = "# 标题\n\n正文内容\n\n## 第二章"
    doc = parser.parse_from_bytes("readme.md", content.encode("utf-8"))
    assert doc.document_id.startswith("doc_")
    assert doc.file_name == "readme.md"
    assert "标题" in doc.pages[0].text


def test_parse_rejects_unsupported():
    parser = FileParser()
    with pytest.raises(FileParserError):
        parser.parse_from_bytes("data.bin", b"binary content")


def test_parse_rejects_fake_extension():
    parser = FileParser()
    with pytest.raises(FileParserError):
        parser.parse_from_bytes("fake.pdf", b"This is not a PDF file")


# ── DOCX 解析 ─────────────────────────────────────────────

def test_parse_docx_basic():
    """创建最小 DOCX 文件测试解析"""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx 未安装")

    parser = FileParser()
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc = Document()
        doc.add_paragraph("甲方：测试公司")
        doc.add_paragraph("乙方：测试员工")
        doc.save(tmp.name)
        tmp_path = Path(tmp.name)

    try:
        result = parser.parse(str(tmp_path))
        assert result.file_name.endswith(".docx")
        assert any("测试公司" in p.text for p in result.pages)
    finally:
        tmp_path.unlink()


# ── MinerU 集成测试（需 conda 环境，CI 跳过）───────────────

@pytest.mark.skip(reason="需要 MinerU conda 环境和 PDF 文件，跳过")
def test_mineru_pdf_parse():
    parser = FileParser()
    pdf_path = Path("data/test/sample.pdf")
    if not pdf_path.exists():
        pytest.skip("测试 PDF 文件不存在")
    doc = parser.parse(str(pdf_path))
    assert doc.document_id.startswith("doc_")
    assert len(doc.pages) > 0
